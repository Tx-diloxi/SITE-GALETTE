from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

# Couleurs
VERT = RGBColor(0x2D, 0x6A, 0x4F)
VERT_CLAIR = RGBColor(0x4A, 0x7C, 0x59)
ROUGE = RGBColor(0xC0, 0x39, 0x2B)
BEIGE = RGBColor(0xF5, 0xF0, 0xE8)
GRIS_FONCE = RGBColor(0x2C, 0x3E, 0x50)
GRIS = RGBColor(0x7F, 0x8C, 0x8D)
GRIS_CLAIR = RGBColor(0xEC, 0xF0, 0xF1)
BLANC = RGBColor(0xFF, 0xFF, 0xFF)
NOIR = RGBColor(0x00, 0x00, 0x00)

SWOT_COLORS = {
    'S': RGBColor(0x27, 0xAE, 0x60),
    'W': RGBColor(0xE7, 0x4C, 0x3C),
    'O': RGBColor(0x29, 0x80, 0xB9),
    'T': RGBColor(0xF3, 0x9C, 0x12),
}

def add_footer(slide, num, total=14):
    left = Inches(0.8)
    top = Inches(7.0)
    width = Inches(11.7)
    height = Inches(0.4)
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    p = tf.paragraphs[0]
    p.text = f"LE SECH Marceau | 26/06/2026 | Slide {num}/{total}"
    p.font.size = Pt(9)
    p.font.color.rgb = GRIS
    p.alignment = PP_ALIGN.CENTER

def add_bg_rect(slide, color=VERT):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0), Inches(0),
        prs.slide_width, Inches(0.15)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()

def add_section_bar(slide, title=None, subtitle=None, num=None):
    add_bg_rect(slide)
    if num:
        add_footer(slide, num)

def new_slide(prs):
    layout = prs.slide_layouts[6]  # blank layout
    slide = prs.slides.add_slide(layout)
    return slide

def add_title_text(slide, text, left=0.8, top=0.4, width=11.7, height=1.0, size=32, bold=True, color=VERT):
    txBox = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.color.rgb = color
    return txBox

def add_body_text(slide, text, left=0.8, top=1.6, width=11.7, height=5.0, size=16, color=NOIR, bold=False, align=PP_ALIGN.LEFT):
    txBox = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(size)
    p.font.color.rgb = color
    p.font.bold = bold
    p.alignment = align
    return tf

def add_bullet_list(slide, items, left=0.8, top=1.8, width=11.7, height=5.0, size=15):
    txBox = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = txBox.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.text = item
        p.font.size = Pt(size)
        p.font.color.rgb = NOIR
        p.space_after = Pt(6)
        p.level = 0
    return tf

def add_placeholder_box(slide, left, top, width, height, label="Ajouter capture ici"):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(left), Inches(top),
        Inches(width), Inches(height)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = GRIS_CLAIR
    shape.line.color.rgb = GRIS
    shape.line.width = Pt(1.5)
    tf = shape.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = label
    p.font.size = Pt(12)
    p.font.color.rgb = GRIS
    p.alignment = PP_ALIGN.CENTER
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    shape.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER

def add_card(slide, title, content, left, top, width=3.5, height=2.5, title_color=VERT_CLAIR):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top),
        Inches(width), Inches(height)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = BLANC
    shape.line.color.rgb = GRIS_CLAIR
    shape.line.width = Pt(1)
    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.2)
    tf.margin_right = Inches(0.2)
    tf.margin_top = Inches(0.15)
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = title_color
    p2 = tf.add_paragraph()
    p2.text = content
    p2.font.size = Pt(11)
    p2.font.color.rgb = NOIR
    p2.space_before = Pt(4)
    return shape

# ============================================================
# SLIDE 1 — Page de garde
# ============================================================
slide = new_slide(prs)
# Bande verte large en haut
shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), prs.slide_width, Inches(2.8))
shape.fill.solid()
shape.fill.fore_color.rgb = VERT
shape.line.fill.background()

# Titre principal
add_title_text(slide, "Refonte complète du site web vitrine", 1.5, 3.2, 10.3, 1.2, 34, True, VERT)
# Sous-titre
add_body_text(slide, "J'aime la Galette — Groupe agroalimentaire breton", 1.5, 4.3, 10.3, 0.6, 20, GRIS_FONCE, False)

# Infos en bas
info_text = "LE SECH Marceau\nIUT de Lannion — Département Informatique — BUT2\nStage du 04/05/2026 au 03/07/2026\nSoutenance : 26 juin 2026 à 11h30"
add_body_text(slide, info_text, 1.5, 5.2, 10.3, 2.0, 14, GRIS_FONCE, False)

# Cercles décoratifs (logo placeholders)
circle = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(0.8), Inches(0.4), Inches(1.8), Inches(1.8))
circle.fill.solid()
circle.fill.fore_color.rgb = BLANC
circle.line.fill.background()
tf = circle.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "Logo\nJ'aime la\nGalette"
p.font.size = Pt(11)
p.font.color.rgb = VERT
p.alignment = PP_ALIGN.CENTER

circle2 = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(10.7), Inches(0.4), Inches(1.8), Inches(1.8))
circle2.fill.solid()
circle2.fill.fore_color.rgb = BLANC
circle2.line.fill.background()
tf2 = circle2.text_frame
tf2.word_wrap = True
p2 = tf2.paragraphs[0]
p2.text = "Logo\nIUT\nLannion"
p2.font.size = Pt(11)
p2.font.color.rgb = VERT
p2.alignment = PP_ALIGN.CENTER

# ============================================================
# SLIDE 2 — Sommaire
# ============================================================
slide = new_slide(prs)
add_section_bar(slide, num=2)
add_title_text(slide, "Sommaire", size=36)

items = [
    "1. Structure d'accueil — J'aime la Galette",
    "2. Sujet du stage — Refonte du site vitrine",
    "3. Organisation du travail — Cycle Analyse / Maquette / Dev / Test",
    "4. Réalisations — Architecture, Backoffice, Chatbot, Visite Mystère, SEO",
    "5. Conclusion — Bilan, PPP, Remerciements",
]
txBox = slide.shapes.add_textbox(Inches(1.5), Inches(2.0), Inches(10), Inches(4.5))
tf = txBox.text_frame
tf.word_wrap = True
for i, item in enumerate(items):
    if i == 0:
        p = tf.paragraphs[0]
    else:
        p = tf.add_paragraph()
    p.text = item
    p.font.size = Pt(18)
    p.font.color.rgb = NOIR
    p.space_after = Pt(16)
    p.level = 0

# ============================================================
# SLIDE 3 — Structure d'accueil (1/2)
# ============================================================
slide = new_slide(prs)
add_section_bar(slide, num=3)
add_title_text(slide, "Structure d'accueil — J'aime la Galette", size=28, top=0.3, height=0.8)

add_body_text(slide, "Groupe agroalimentaire breton — Fabricant artisanal de crêpes et galettes", 0.8, 1.1, 11.7, 0.5, 16, GRIS, False)

# Cards for key info
add_card(slide, "🏭 2 marques", "• J'aime la Galette (historique)\n• Be Good'n (gourmand/responsable)", 0.8, 1.8, 3.5, 1.5)
add_card(slide, "📍 8 ateliers + 1 siège", "• Broons, Alençon, Angers…\n• Bretagne et Pays de la Loire", 4.7, 1.8, 3.5, 1.5)
add_card(slide, "👥 Dirigeant", "• Pascal ENAULT — Tuteur entreprise\n• Référent direct du stage", 8.6, 1.8, 3.9, 1.5)

add_card(slide, "📦 Produits", "• Galettes, crêpes, chips\n• Caramel, cidre\n• Distribution GMS + RHD", 0.8, 3.6, 3.5, 1.5)
add_card(slide, "🎯 Publics cibles", "• Consommateurs finaux\n• Distributeurs, Collectivités\n• Candidats RH, Presse", 4.7, 3.6, 3.5, 1.5)
add_card(slide, "🔄 Relations", "• Équipe : direction commerciale\n• Utilisateurs : visiteurs, admins\n       commerciaux, livreurs", 8.6, 3.6, 3.9, 1.5)

add_body_text(slide, "Contexte : stage BUT2 du 04/05/2026 au 03/07/2026 (9 semaines)", 0.8, 5.5, 11.7, 0.4, 12, GRIS, False)

# ============================================================
# SLIDE 4 — Structure d'accueil (2/2) SWOT
# ============================================================
slide = new_slide(prs)
add_section_bar(slide, num=4)
add_title_text(slide, "Analyse SWOT & Problématique", size=28, top=0.3, height=0.8)

# SWOT boxes - 2 columns, 2 rows
swot_data = [
    ("S — Forces", SWOT_COLORS['S'],
     ["Produits artisanaux bretons de qualité",
      "Savoir-faire industriel + valeurs RSE fortes"]),
    ("W — Faiblesses", SWOT_COLORS['W'],
     ["Site web obsolète, absence de backoffice",
      "Visibilité en ligne limitée (SEO inexistant)"]),
    ("O — Opportunités", SWOT_COLORS['O'],
     ["Boom du snacking sain et local",
      "SEO/LLMO pour capter une nouvelle clientèle"]),
    ("T — Menaces", SWOT_COLORS['T'],
     ["Concurrence des grandes industries agro",
      "Évolution rapide des attentes web mobiles"]),
]

positions = [
    (0.8, 1.3, 5.5, 2.0),
    (7.0, 1.3, 5.5, 2.0),
    (0.8, 3.5, 5.5, 2.0),
    (7.0, 3.5, 5.5, 2.0),
]

for (title, color, items), (l, t, w, h) in zip(swot_data, positions):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(l), Inches(t), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = BLANC
    shape.line.color.rgb = color
    shape.line.width = Pt(2)
    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.2)
    tf.margin_right = Inches(0.2)
    tf.margin_top = Inches(0.1)
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = color
    for item in items:
        p2 = tf.add_paragraph()
        p2.text = "• " + item
        p2.font.size = Pt(11)
        p2.font.color.rgb = NOIR
        p2.space_before = Pt(3)

# Problématique
shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(5.7), Inches(11.7), Inches(1.0))
shape.fill.solid()
shape.fill.fore_color.rgb = VERT
shape.line.fill.background()
tf = shape.text_frame
tf.word_wrap = True
tf.margin_left = Inches(0.3)
p = tf.paragraphs[0]
p.text = "❓ Problématique : Comment moderniser la présence en ligne du groupe pour mieux vendre, recruter et communiquer ?"
p.font.size = Pt(14)
p.font.bold = True
p.font.color.rgb = BLANC
p.alignment = PP_ALIGN.CENTER

# ============================================================
# SLIDE 5 — Sujet du stage
# ============================================================
slide = new_slide(prs)
add_section_bar(slide, num=5)
add_title_text(slide, "Sujet du stage — Objectifs & Missions", size=28, top=0.3, height=0.8)

add_body_text(slide, "Objectif : Refonte complète du site vitrine — site dynamique, backoffice, SEO", 0.8, 1.1, 11.7, 0.5, 16, GRIS, False)

# Cycle: Analyse → Maquette → Dev → Test
cycle_data = [
    ("Analyse", "Besoin client\nEntretiens Pascal ENAULT\nCahier des charges"),
    ("Maquette", "Maquettage Figma\nValidation visuelle\navant développement"),
    ("Développement", "PHP procédural + PDO\n53 tables MySQL\nFront controller"),
    ("Test", "Tests fonctionnels\nValidation responsive\nRecette client"),
]
for i, (title, content) in enumerate(cycle_data):
    l = 0.8 + i * 3.1
    t = 1.8
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(l), Inches(t), Inches(2.8), Inches(2.2))
    shape.fill.solid()
    shape.fill.fore_color.rgb = VERT_CLAIR if i % 2 == 0 else BLANC
    shape.line.color.rgb = VERT_CLAIR
    shape.line.width = Pt(1.5)
    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.15)
    tf.margin_right = Inches(0.15)
    tf.margin_top = Inches(0.1)
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = BLANC if i % 2 == 0 else VERT_CLAIR
    p.alignment = PP_ALIGN.CENTER
    p2 = tf.add_paragraph()
    p2.text = content
    p2.font.size = Pt(11)
    p2.font.color.rgb = BLANC if i % 2 == 0 else NOIR
    p2.alignment = PP_ALIGN.CENTER
    p2.space_before = Pt(8)

# Missions list
missions = [
    "1. Architecture BDD (53 tables) + Routage (front controller, URLs propres)",
    "2. 10 pages publiques avec contenu dynamique (29 partials)",
    "3. Backoffice CRUD complet (20 formulaires de sections, produits, marques, FAQ…)",
    "4. Chatbot intelligent (FULLTEXT + scoring), SEO (JSON-LD, sitemap), RGPD",
    "5. Module Visite Mystère pour commerciaux (notation, photos)",
]
add_bullet_list(slide, missions, 0.8, 4.3, 11.7, 3.0, 13)

# ============================================================
# SLIDE 6 — Planning Gantt
# ============================================================
slide = new_slide(prs)
add_section_bar(slide, num=6)
add_title_text(slide, "Planning du stage — 9 semaines", size=28, top=0.3, height=0.8)

# Table-like Gantt using text boxes
gantt_data = [
    ("Analyse & Conception", "Sem 1-2", "Analyse besoin, maquettage Figma, conception BDD + routage", VERT),
    ("Développement pages", "Sem 3-5", "10 pages publiques, 29 partials, intégration CSS/JS", VERT_CLAIR),
    ("Backoffice", "Sem 6-7", "Dashboard, CRUD pages/produits/marques/sites/FAQ, auth", VERT_CLAIR),
    ("Chatbot & SEO", "Sem 8", "Algorithme FULLTEXT, JSON-LD, sitemap, meta SEO", VERT_CLAIR),
    ("Visite Mystère", "Sem 8", "Module inspection, notation, upload photos", VERT_CLAIR),
    ("Tests & Recette", "Sem 9", "Tests fonctionnels, corrections, validation client", VERT),
]

for i, (phase, period, desc, color) in enumerate(gantt_data):
    t = 1.3 + i * 0.85
    # Barre colorée
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(t), Inches(11.7), Inches(0.65))
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.2)
    tf.margin_top = Inches(0.05)
    p = tf.paragraphs[0]
    p.text = f"{phase}  |  {period}"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = BLANC
    p2 = tf.add_paragraph()
    p2.text = desc
    p2.font.size = Pt(11)
    p2.font.color.rgb = BLANC
    p2.space_before = Pt(2)

# Milestone
shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(3.5), Inches(6.7), Inches(6.3), Inches(0.5))
shape.fill.solid()
shape.fill.fore_color.rgb = ROUGE
shape.line.fill.background()
tf = shape.text_frame
p = tf.paragraphs[0]
p.text = "🔴 Jalon : Mise en ligne — Site complet + backoffice opérationnel"
p.font.size = Pt(13)
p.font.bold = True
p.font.color.rgb = BLANC
p.alignment = PP_ALIGN.CENTER

# ============================================================
# SLIDE 7 — Organisation du travail
# ============================================================
slide = new_slide(prs)
add_section_bar(slide, num=7)
add_title_text(slide, "Organisation du travail", size=28, top=0.3, height=0.8)

# Méthode
add_card(slide, "📋 Méthodologie",
         "Cycle itératif :\nAnalyse → Maquette → Dev → Test\nInspiré des méthodes agiles\nAjustements continus avec le tuteur", 0.8, 1.3, 3.8, 2.5)

add_card(slide, "💻 Environnement technique",
         "Docker (PHP 8.2 + MySQL 8.0\n+ phpMyAdmin)\nVS Code, Git\nTesté sur http://localhost:8080", 4.9, 1.3, 3.8, 2.5)

add_card(slide, "🔧 Stack",
         "PHP procédural + PDO\nHTML/CSS vanilla + SCSS\nJS vanilla, PHPMailer\nComposer (1 dépendance)", 9.0, 1.3, 3.5, 2.5)

# Justification
shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(4.2), Inches(11.7), Inches(2.5))
shape.fill.solid()
shape.fill.fore_color.rgb = BEIGE
shape.line.color.rgb = VERT_CLAIR
shape.line.width = Pt(1)
tf = shape.text_frame
tf.word_wrap = True
tf.margin_left = Inches(0.3)
tf.margin_top = Inches(0.15)
p = tf.paragraphs[0]
p.text = "Pourquoi ce choix technique ?"
p.font.size = Pt(15)
p.font.bold = True
p.font.color.rgb = VERT
p2 = tf.add_paragraph()
p2.text = "→ Contrainte forte : hébergement mutualisé PHP/MySQL standard (pas de framework, pas de Node, pas d'ORM)"
p2.font.size = Pt(12)
p2.font.color.rgb = NOIR
p2.space_before = Pt(6)
p3 = tf.add_paragraph()
p3.text = "→ PHP procédural + PDO = compatible avec tous les hébergeurs mutualisés"
p3.font.size = Pt(12)
p3.font.color.rgb = NOIR
p4 = tf.add_paragraph()
p4.text = "→ SCSS compilé en CSS = pas de build complexe en production"
p4.font.size = Pt(12)
p4.font.color.rgb = NOIR

# ============================================================
# SLIDE 8 — Réalisations : Maquette → Code
# ============================================================
slide = new_slide(prs)
add_section_bar(slide, num=8)
add_title_text(slide, "Réalisation — De la maquette au code", size=28, top=0.3, height=0.8)

add_body_text(slide, "Processus complet pour chaque page : maquettage Figma → développement → test", 0.8, 1.1, 11.7, 0.5, 16, GRIS, False)

# Left: Maquette placeholder
add_placeholder_box(slide, 0.8, 1.8, 5.5, 3.5, "Placeholder :\nCapture maquette Figma")
add_title_text(slide, "Maquette Figma", 0.8, 5.5, 5.5, 0.4, 14, True, VERT)

# Right: Code placeholder
add_placeholder_box(slide, 7.0, 1.8, 5.5, 3.5, "Placeholder :\nCapture page réalisée")
add_title_text(slide, "Page finale", 7.0, 5.5, 5.5, 0.4, 14, True, VERT)

# Arrow between
shape = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(6.0), Inches(3.0), Inches(0.8), Inches(0.6))
shape.fill.solid()
shape.fill.fore_color.rgb = VERT
shape.line.fill.background()

# Bottom text
add_body_text(slide, "✅ Validation client à chaque étape — itérations rapides (1-2 jours par page)", 0.8, 6.2, 11.7, 0.4, 12, GRIS, False)

# ============================================================
# SLIDE 9 — Architecture & BDD
# ============================================================
slide = new_slide(prs)
add_section_bar(slide, num=9)
add_title_text(slide, "Réalisation — Architecture & Base de données", size=28, top=0.3, height=0.8)

# Architecture
add_card(slide, "🔀 Front Controller",
         "public/index.php lit l'URL\n→ .htaccess rewrite\n→ mapping vers pages/*.php\n→ include de partials", 0.8, 1.3, 3.8, 2.2)

add_card(slide, "🗄️ Base de données",
         "53 tables MySQL (InnoDB)\nModèle relationnel\nClés étrangères CASCADE\nIndex FULLTEXT (FAQ)", 4.9, 1.3, 3.8, 2.2)

add_card(slide, "📄 Architecture page",
         "1. Load data (PDO queries)\n2. Include head + header\n3. Include partials avec données\n4. Include footer + chatbot", 9.0, 1.3, 3.5, 2.2)

# Table listing
add_body_text(slide, "Tables principales :", 0.8, 3.7, 11.7, 0.3, 13, VERT, True)

tables_text = ("PUBLIC : hero, marque, produit, partial_*, card_* (20+ tables de contenu)\n"
               "FORMULAIRES : formulaire_contact, applications, candidature_spontanee\n"
               "FAQ : categorie_FAQ, question_FAQ, chatbot_log\n"
               "VISITE MYSTÈRE : commercial, livreur, inspection, inspection_reponse, inspection_photo")
add_body_text(slide, tables_text, 0.8, 4.1, 11.7, 1.2, 11, NOIR, False)

# Code snippet placeholder
add_placeholder_box(slide, 0.8, 5.3, 11.7, 1.5, "Placeholder : extrait de code du routage (index.php) + schéma relationnel BDD")

# ============================================================
# SLIDE 10 — Backoffice
# ============================================================
slide = new_slide(prs)
add_section_bar(slide, num=10)
add_title_text(slide, "Réalisation — Backoffice complet", size=28, top=0.3, height=0.8)

add_card(slide, "📊 Dashboard",
         "Indicateurs : candidatures non lues,\nmessages non lus, produits en ligne,\noffres actives, inspections\nListe des dernières activités", 0.8, 1.3, 3.8, 2.2)

add_card(slide, "📝 Éditeur de pages",
         "20 formulaires de section\nÉdition de tout le contenu\nsans toucher au code\nUpload d'images (5 Mo max)", 4.9, 1.3, 3.8, 2.2)

add_card(slide, "🔐 Authentification",
         "Session PHP, timeout 2h\nCSRF tokens sur tous les POST\nBlocage 15 min après 5 échecs\nDouble rôle : admin + commercial", 9.0, 1.3, 3.5, 2.2)

add_card(slide, "📦 CRUD",
         "Produits, Marques, Sites, FAQ,\nCommerciaux, Livreurs\nFiltres, pagination, tri\nToggle en ligne/hors ligne", 0.8, 3.8, 3.8, 2.0)

add_card(slide, "📬 Gestion contacts",
         "Liste des messages avec filtres\n(statut + profil B2C/B2B)\nMarquage lu/traité/archivé\nNote admin", 4.9, 3.8, 3.8, 2.0)

add_card(slide, "📄 Candidatures",
         "Consultation des CV uploadés\nVisualisation des candidatures\nspontanées et offres\nMarquage lu/non lu", 9.0, 3.8, 3.5, 2.0)

add_placeholder_box(slide, 0.8, 6.0, 11.7, 0.9, "Placeholder : capture d'écran du dashboard admin + formulaire d'édition")

# ============================================================
# SLIDE 11 — Chatbot
# ============================================================
slide = new_slide(prs)
add_section_bar(slide, num=11)
add_title_text(slide, "Réalisation — Chatbot intelligent", size=28, top=0.3, height=0.8)

add_body_text(slide, "Widget flottant sur toutes les pages — interroge la FAQ en temps réel", 0.8, 1.1, 11.7, 0.5, 16, GRIS, False)

# Algorithme steps
steps = [
    ("Étape 1", "FULLTEXT BOOLEAN MODE", "Recherche avec tous les tokens (opérateur +)\nIndex FULLTEXT sur question + mots_cles", VERT),
    ("Étape 2", "Fallback LIKE", "Si aucun résultat FULLTEXT → LIKE sur chaque token\nMatch partiel acceptable", VERT_CLAIR),
    ("Étape 3", "Scoring", "Ratio de tokens matchés\nBonus si match sur question (vs mots-clés)\nSeuil : 0.3 FULLTEXT, 0.7 LIKE", VERT_CLAIR),
    ("Étape 4", "Réponse", "Renvoi : réponse, catégorie, URL FAQ\n + suggestions de questions\nLog dans chatbot_log", VERT),
]

for i, (title, sub, desc, color) in enumerate(steps):
    l = 0.8 + i * 3.1
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(l), Inches(1.8), Inches(2.8), Inches(2.5))
    shape.fill.solid()
    shape.fill.fore_color.rgb = BLANC
    shape.line.color.rgb = color
    shape.line.width = Pt(2)
    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.15)
    tf.margin_right = Inches(0.15)
    tf.margin_top = Inches(0.1)
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = color
    p.alignment = PP_ALIGN.CENTER
    p2 = tf.add_paragraph()
    p2.text = sub
    p2.font.size = Pt(11)
    p2.font.bold = True
    p2.font.color.rgb = GRIS_FONCE
    p2.alignment = PP_ALIGN.CENTER
    p2.space_before = Pt(4)
    p3 = tf.add_paragraph()
    p3.text = desc
    p3.font.size = Pt(10)
    p3.font.color.rgb = NOIR
    p3.space_before = Pt(6)

add_card(slide, "📊 Stats FAQ",
         "• 3 catégories, 9 questions\n• Logging de toutes les questions\n• Amélioration continue du matching", 0.8, 4.6, 5.5, 1.5)

add_card(slide, "🌐 LLMO/GEO",
         "• H2 formulés en questions\n• Contenu factuel et chiffré\n• Blocs 'En résumé' pour les LLM", 7.0, 4.6, 5.5, 1.5)

add_placeholder_box(slide, 0.8, 6.3, 11.7, 0.6, "Placeholder : capture du widget chatbot + exemple de réponse")

# ============================================================
# SLIDE 12 — Visite Mystère
# ============================================================
slide = new_slide(prs)
add_section_bar(slide, num=12)
add_title_text(slide, "Réalisation — Module Visite Mystère", size=28, top=0.3, height=0.8)

add_body_text(slide, "Application métier pour les commerciaux — inspection des livreurs sur le terrain", 0.8, 1.1, 11.7, 0.5, 16, GRIS, False)

add_card(slide, "🔑 Espace commercial",
         "Authentification dédiée\n(utilisateurs table commercial)\nSession séparée de l'admin", 0.8, 1.8, 3.5, 2.2)

add_card(slide, "📋 Grille d'inspection",
         "10 questions, 3 catégories :\n• Propreté/Hygiène\n• Merchandising\n• Professionnalisme\nNotation 1 à 5 étoiles", 4.7, 1.8, 3.5, 2.2)

add_card(slide, "📸 Upload photos",
         "Drag & drop jusqu'à 10 photos\nAperçu avant validation\nStockage par inspection\n(Dossier dédié)", 8.6, 1.8, 3.9, 2.2)

add_card(slide, "📊 Score /20",
         "Calcul automatique\nCommentaire général\nHistorique consultable\n4 tables liées", 0.8, 4.3, 5.5, 1.5)

add_card(slide, "👤 Utilisateurs",
         "• 2 commerciaux\n• 5 livreurs\nDonnées en seed SQL", 6.7, 4.3, 5.8, 1.5)

add_placeholder_box(slide, 0.8, 6.1, 11.7, 0.8, "Placeholder : capture du formulaire d'inspection avec notation étoiles")

# ============================================================
# SLIDE 13 — SEO & RGPD
# ============================================================
slide = new_slide(prs)
add_section_bar(slide, num=13)
add_title_text(slide, "Réalisation — SEO, Formulaires & RGPD", size=28, top=0.3, height=0.8)

# SEO
shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.3), Inches(5.7), Inches(2.3))
shape.fill.solid()
shape.fill.fore_color.rgb = BLANC
shape.line.color.rgb = VERT
shape.line.width = Pt(1.5)
tf = shape.text_frame
tf.word_wrap = True
tf.margin_left = Inches(0.2)
tf.margin_top = Inches(0.1)
p = tf.paragraphs[0]
p.text = "🔍 Optimisation SEO"
p.font.size = Pt(14)
p.font.bold = True
p.font.color.rgb = VERT
items = [
    "Titres uniques + meta descriptions 150-160 car.",
    "JSON-LD : Organization, BreadcrumbList, FAQPage, Product, LocalBusiness",
    "Sitemap.xml dynamique (10 pages)",
    "robots.txt, balises OG + Twitter Cards",
    "H2 formulés en questions (LLMO/GEO)",
    "Images optimisées, lazy loading",
]
for item in items:
    p2 = tf.add_paragraph()
    p2.text = "• " + item
    p2.font.size = Pt(11)
    p2.font.color.rgb = NOIR
    p2.space_before = Pt(3)

# Formulaires & RGPD
shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.3), Inches(5.7), Inches(2.3))
shape.fill.solid()
shape.fill.fore_color.rgb = BLANC
shape.line.color.rgb = VERT
shape.line.width = Pt(1.5)
tf = shape.text_frame
tf.word_wrap = True
tf.margin_left = Inches(0.2)
tf.margin_top = Inches(0.1)
p = tf.paragraphs[0]
p.text = "📝 Formulaires & RGPD"
p.font.size = Pt(14)
p.font.bold = True
p.font.color.rgb = VERT
items = [
    "Contact : double profil B2C/B2B + SMTP PHPMailer",
    "Recrutement : CV upload + candidature spontanée",
    "CSRF token sur tous les formulaires POST",
    "Case consentement RGPD obligatoire",
    "Protection MIME + taille upload (5 Mo)",
    "Validation serveur + client",
]
for item in items:
    p2 = tf.add_paragraph()
    p2.text = "• " + item
    p2.font.size = Pt(11)
    p2.font.color.rgb = NOIR
    p2.space_before = Pt(3)

# Stats / KPI
shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(3.9), Inches(11.7), Inches(1.5))
shape.fill.solid()
shape.fill.fore_color.rgb = BEIGE
shape.line.fill.background()
tf = shape.text_frame
tf.word_wrap = True
tf.margin_left = Inches(0.3)
tf.margin_top = Inches(0.1)
p = tf.paragraphs[0]
p.text = "Chiffres clés du projet"
p.font.size = Pt(14)
p.font.bold = True
p.font.color.rgb = VERT

kpi = "10 pages publiques  |  29 partials  |  53 tables MySQL  |  20+ formulaires admin\n9 questions FAQ  |  10 questions inspection  |  5 types d'utilisateurs"
p2 = tf.add_paragraph()
p2.text = kpi
p2.font.size = Pt(13)
p2.font.color.rgb = NOIR
p2.space_before = Pt(6)

add_placeholder_box(slide, 0.8, 5.7, 5.5, 1.2, "Placeholder : capture Lighthouse / SEO score")

add_placeholder_box(slide, 7.0, 5.7, 5.5, 1.2, "Placeholder : capture formulaire contact")

# ============================================================
# SLIDE 14 — Conclusion
# ============================================================
slide = new_slide(prs)
add_section_bar(slide, num=14)
add_title_text(slide, "Conclusion", size=36, top=0.3, height=0.8)

# Bilan
add_card(slide, "✅ Bilan",
         "Objectifs atteints :\n• Site vitrine dynamique complet\n• Backoffice CRUD opérationnel\n• Chatbot intelligent + SEO optimisé\n• Module Visite Mystère fonctionnel", 0.8, 1.3, 3.8, 2.5, VERT)

add_card(slide, "📋 Travail restant",
         "• Pages légales (mentions, cookies)\n• Routes produits dynamiques\n• Optimisations mobiles continues\n• Contenus à enrichir", 4.9, 1.3, 3.8, 2.5, ROUGE)

add_card(slide, "🎯 Projet Personnel Pro.",
         "Ce stage conforte mon orientation\nvers le développement web full-stack.\nJ'ai découvert :\n• L'importance du maquettage\n• Le cycle analyse/dev/test en réel\n• Les contraintes d'hébergement", 9.0, 1.3, 3.5, 2.5, VERT)

# Remerciements
shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(4.2), Inches(11.7), Inches(2.5))
shape.fill.solid()
shape.fill.fore_color.rgb = VERT
shape.line.fill.background()
tf = shape.text_frame
tf.word_wrap = True
tf.margin_left = Inches(0.3)
tf.margin_top = Inches(0.15)
p = tf.paragraphs[0]
p.text = "Remerciements"
p.font.size = Pt(18)
p.font.bold = True
p.font.color.rgb = BLANC
p.alignment = PP_ALIGN.CENTER

remerciements = [
    "Pascal ENAULT — Dirigeant J'aime la Galette, pour sa confiance et son accompagnement",
    "Yannick Favreau & Tiphaine Jezequel — Enseignants référents IUT Lannion",
    "L'équipe pédagogique du Département Informatique — IUT de Lannion",
]
for r in remerciements:
    p2 = tf.add_paragraph()
    p2.text = "• " + r
    p2.font.size = Pt(13)
    p2.font.color.rgb = BLANC
    p2.space_before = Pt(6)
    p2.alignment = PP_ALIGN.LEFT

# Question ouverte
p3 = tf.add_paragraph()
p3.text = "\nDes questions ?"
p3.font.size = Pt(18)
p3.font.bold = True
p3.font.color.rgb = BLANC
p3.alignment = PP_ALIGN.CENTER
p3.space_before = Pt(12)

# ============================================================
# SAUVEGARDE
# ============================================================
output_path = "/home/etudiant/opencode/SITE GALETTE/Soutenance_Stage_LE_SECH_Marceau.pptx"
prs.save(output_path)
print(f"✅ Présentation générée : {output_path}")
print(f"📊 {len(prs.slides)} slides créées")
